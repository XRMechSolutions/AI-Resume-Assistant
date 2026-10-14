import os
import sys
import re
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

MAX_PAGES = 2

# Canonical list of recognized section headers.
# Used by both the parser (to detect sections) and the renderer (to order output).
SECTION_PATTERNS = [
    'SUMMARY', 'PROFESSIONAL SUMMARY', 'OBJECTIVE',
    'TECHNICAL SKILLS', 'SKILLS', 'CORE COMPETENCIES', 'CORE TECHNICAL CAPABILITIES',
    'PROFESSIONAL EXPERIENCE', 'WORK EXPERIENCE', 'EXPERIENCE',
    'EDUCATION',
    'NOTABLE PROJECTS', 'KEY PROJECTS', 'PROJECTS',
    'KEY ACHIEVEMENTS', 'ACHIEVEMENTS',
    'CERTIFICATIONS', 'CERTIFICATIONS & TRAINING',
    'ADDITIONAL INFORMATION', 'PERSONAL INTERESTS'
]

# Spanish headers for Spanish-language resumes. Each maps to a canonical section above,
# and the original header text is kept so it renders in Spanish.
SECTION_ALIASES = {
    'PERFIL PROFESIONAL': 'SUMMARY',
    'RESUMEN': 'SUMMARY',
    'COMPETENCIAS TÉCNICAS': 'TECHNICAL SKILLS',
    'HABILIDADES TÉCNICAS': 'TECHNICAL SKILLS',
    'EXPERIENCIA PROFESIONAL': 'PROFESSIONAL EXPERIENCE',
    'FORMACIÓN': 'EDUCATION',
    'FORMACIÓN ACADÉMICA': 'EDUCATION',
    'PROYECTOS DESTACADOS': 'KEY PROJECTS',
    'LOGROS DESTACADOS': 'KEY ACHIEVEMENTS',
    'IDIOMAS': 'ADDITIONAL INFORMATION',
}

# Month abbreviations accepted in job date ranges (English, then Spanish)
_MONTHS = r'(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|Ene|Abr|Ago|Dic)'
_DATE_RANGE = re.compile(
    rf'\b({_MONTHS}\s+\d{{4}}\s*[-–]\s*(?:Present|Actualidad|\d{{4}}|{_MONTHS}\s+\d{{4}}))\b'
)


def _esc(text):
    """Escape ampersands for ReportLab XML markup."""
    return text.replace('&', '&amp;')


def parse_resume_txt(file_path, verbose=False):
    """Parse a resume text file and extract sections."""
    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()

    sections = {}
    titles = {}  # section key -> display title, only for aliased (non-English) headers
    current_section = None
    current_content = []
    warnings = []

    lines = content.split('\n')
    
    # Extract header info (name, title, contact) - look for contact info patterns
    header_lines = []
    i = 0
    header_complete = False
    
    while i < len(lines) and not header_complete:
        line = lines[i].strip()
        
        # Check if we've hit a major section header
        if line and (any(line.upper() == pattern for pattern in ['PROFESSIONAL SUMMARY', 'SUMMARY', 'OBJECTIVE', 'TECHNICAL SKILLS', 'SKILLS'])
                     or SECTION_ALIASES.get(line.upper()) in ('SUMMARY', 'TECHNICAL SKILLS')):
            header_complete = True
            break
        
        # Add non-empty lines to header until we hit a section
        if line:
            header_lines.append(line)
        elif header_lines and len(header_lines) >= 2:  # Empty line after we have at least name and one contact line
            header_complete = True
            break
        
        i += 1
    
    sections['header'] = header_lines
    
    section_patterns = SECTION_PATTERNS
    header_name = header_lines[0].strip().upper() if header_lines else ''
    
    # Process the rest of the file from where header ended
    for line_idx in range(i, len(lines)):
        line = lines[line_idx].rstrip()
        
        # Check if this line is a section header
        is_section = False
        alias = SECTION_ALIASES.get(line.upper().strip())
        for pattern in section_patterns:
            if line.upper().strip() == pattern or alias == pattern:
                # Save previous section
                if current_section:
                    sections[current_section] = '\n'.join(current_content).strip()
                # Start new section
                current_section = pattern.replace(' ', '_').upper()
                current_content = []
                if alias:
                    titles[current_section] = line.strip().upper()
                is_section = True
                break
        
        if not is_section:
            # Detect potential unrecognized section headers:
            # ALL CAPS, 2+ words, not a bullet, not blank, not a known skill/contact line
            stripped = line.strip()
            if (stripped and stripped == stripped.upper() and len(stripped) >= 4
                    and not stripped.startswith(('•', '-', '('))
                    and ':' not in stripped and '|' not in stripped
                    and not re.search(r'\d{4}', stripped)
                    and len(stripped.split()) >= 2
                    and stripped != header_name):
                warnings.append(f"  [WARN] Possible unrecognized section header: \"{stripped}\" (line {line_idx + 1})")

            if current_section or line.strip():  # Add to current section or capture orphaned content
                current_content.append(line)

    # Save the last section
    if current_section:
        sections[current_section] = '\n'.join(current_content).strip()
    sections['_titles'] = titles

    # Print warnings
    if warnings:
        basename = os.path.basename(file_path)
        print(f"\n  Warnings for {basename}:")
        for w in warnings:
            print(w)

    if verbose:
        basename = os.path.basename(file_path)
        print(f"\n  Parsed sections in {basename}:")
        for key, val in sections.items():
            if key == '_titles':
                continue
            if key == 'header':
                print(f"    HEADER: {len(val)} lines")
            else:
                line_count = len([l for l in val.split('\n') if l.strip()])
                print(f"    {key}: {line_count} content lines")

    return sections

# Arial ships with Windows; elsewhere register_fonts() fails and the PDF falls back to Helvetica
FONT_DIR = os.path.join(os.environ.get('WINDIR', 'C:/Windows'), 'Fonts')

def register_fonts():
    """Register embedded TTF fonts for a properly sized PDF output."""
    try:
        pdfmetrics.registerFont(TTFont('Arial', f'{FONT_DIR}/arial.ttf'))
        pdfmetrics.registerFont(TTFont('Arial-Bold', f'{FONT_DIR}/arialbd.ttf'))
        pdfmetrics.registerFont(TTFont('Arial-Italic', f'{FONT_DIR}/ariali.ttf'))
        pdfmetrics.registerFont(TTFont('Arial-BoldItalic', f'{FONT_DIR}/arialbi.ttf'))
        from reportlab.pdfbase.pdfmetrics import registerFontFamily
        registerFontFamily('Arial', normal='Arial', bold='Arial-Bold',
                           italic='Arial-Italic', boldItalic='Arial-BoldItalic')
        return True
    except Exception:
        return False

def build_resume_elements(txt_file_path, verbose=False):
    """Parse a resume .txt and build the ReportLab flowables shared by the PDF and DOCX outputs."""
    fonts_available = register_fonts()
    base_font = 'Arial' if fonts_available else 'Helvetica'
    base_font_bold = 'Arial-Bold' if fonts_available else 'Helvetica-Bold'

    # Parse the resume
    sections = parse_resume_txt(txt_file_path, verbose=verbose)

    elements = []
    styles = getSampleStyleSheet()

    # Define custom styles
    name_style = ParagraphStyle(
        'CustomName',
        parent=styles['Heading1'],
        fontSize=16,
        fontName=base_font_bold,
        textColor=colors.HexColor('#1a1a1a'),
        spaceAfter=2,
        alignment=TA_CENTER,
        leading=18
    )

    contact_style = ParagraphStyle(
        'ContactInfo',
        parent=styles['Normal'],
        fontSize=10,
        fontName=base_font,
        textColor=colors.HexColor('#444444'),
        spaceAfter=2,
        alignment=TA_CENTER,
        leading=12
    )

    section_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontSize=11,
        textColor=colors.HexColor('#000080'),
        spaceAfter=6,
        spaceBefore=10,
        leading=13,
        fontName=base_font_bold
    )

    normal_style = ParagraphStyle(
        'CustomNormal',
        parent=styles['Normal'],
        fontSize=10,
        fontName=base_font,
        leading=12,
        spaceAfter=6,
        alignment=TA_JUSTIFY
    )

    bullet_style = ParagraphStyle(
        'BulletStyle',
        parent=styles['Normal'],
        fontSize=10,
        fontName=base_font,
        leading=12,
        leftIndent=15,
        spaceAfter=2,
        alignment=TA_LEFT
    )
    
    job_title_style = ParagraphStyle(
        'JobTitle',
        parent=styles['Normal'],
        fontSize=10.5,
        textColor=colors.HexColor('#000000'),
        spaceAfter=3,
        spaceBefore=6,
        leading=12,
        fontName=base_font_bold
    )
    
    company_style = ParagraphStyle(
        'CompanyStyle',
        parent=styles['Normal'],
        fontSize=10,
        fontName=base_font,
        textColor=colors.HexColor('#333333'),
        spaceAfter=2,
        leading=12
    )
    
    # Process header
    if 'header' in sections and sections['header']:
        header_lines = sections['header']
        
        # First line is usually the name - bold the entire name
        if header_lines:
            name = f"<b>{header_lines[0]}</b>"
            elements.append(Paragraph(name, name_style))
        
        # Rest are contact details
        for line in header_lines[1:]:
            if line:
                # Make links clickable
                line = re.sub(r'(github\.com/\S+)', r'<link href="https://\1">\1</link>', line)
                line = re.sub(r'(linkedin\.com/\S+)', r'<link href="https://\1">\1</link>', line)
                line = re.sub(r'(\S+@\S+\.\S+)', r'<link href="mailto:\1">\1</link>', line)
                elements.append(Paragraph(line, contact_style))
        
        elements.append(Spacer(1, 12))

    # Process sections in a logical order
    # Derive section_order from the canonical SECTION_PATTERNS list
    # so the two can never drift out of sync
    section_order = [p.replace(' ', '_').upper() for p in SECTION_PATTERNS]
    
    for section_key in section_order:
        if section_key in sections and sections[section_key]:
            content = sections[section_key]
            
            # Add section heading
            section_title = sections['_titles'].get(section_key, section_key.replace('_', ' '))
            elements.append(Paragraph(f"<b>{section_title}</b>", section_style))
            
            # Process content based on section type
            if section_key in ['PROFESSIONAL_EXPERIENCE', 'WORK_EXPERIENCE', 'EXPERIENCE']:
                # Parse job entries
                process_experience_section(content, elements, job_title_style, company_style, bullet_style)
            
            elif section_key in ['TECHNICAL_SKILLS', 'SKILLS', 'CORE_COMPETENCIES', 'CORE_TECHNICAL_CAPABILITIES']:
                # Parse skills section
                process_skills_section(content, elements, bullet_style, normal_style)
            
            elif section_key in ['NOTABLE_PROJECTS', 'KEY_PROJECTS', 'PROJECTS']:
                # Parse projects
                process_projects_section(content, elements, job_title_style, bullet_style)
            
            elif section_key in ['EDUCATION']:
                # Parse education
                process_education_section(content, elements, normal_style, bullet_style)
            
            elif section_key in ['KEY_ACHIEVEMENTS', 'ACHIEVEMENTS', 'CERTIFICATIONS', 'CERTIFICATIONS_&_TRAINING']:
                # Parse as bullet points
                process_bullet_section(content, elements, bullet_style)
            
            else:
                # Default: treat as paragraph or multi-line content
                paragraphs = content.split('\n\n')  # Split by double newline for paragraphs
                for para in paragraphs:
                    para = para.strip()
                    if para:
                        # Join lines within a paragraph with spaces
                        para_text = ' '.join(line.strip() for line in para.split('\n') if line.strip())
                        if para_text:
                            elements.append(Paragraph(para_text, normal_style))
    
    return sections, elements


def create_resume_pdf(txt_file_path, output_pdf_path=None, verbose=False):
    """Create a PDF from a resume text file."""

    if not os.path.exists(txt_file_path):
        print(f"Error: File {txt_file_path} not found!")
        return

    sections, elements = build_resume_elements(txt_file_path, verbose=verbose)

    # Determine output filename
    if output_pdf_path is None:
        base_name = os.path.splitext(txt_file_path)[0]
        output_pdf_path = base_name + ".pdf"

    # Create the PDF
    doc = SimpleDocTemplate(output_pdf_path, pagesize=letter,
                            rightMargin=0.5*inch, leftMargin=0.5*inch,
                            topMargin=0.5*inch, bottomMargin=0.5*inch)

    # Embed metadata in PDF Info dictionary
    name = sections['header'][0].strip().title() if sections.get('header') else 'Resume'
    doc.title = f"{name} - Resume"
    doc.author = name
    doc.subject = "Professional Resume"
    doc.creator = "AI Resume Assistant (create_pdf.py)"
    doc.keywords = "resume"

    # Build PDF
    try:
        doc.build(elements)
        print(f"[SUCCESS] Resume PDF created: {output_pdf_path} ({doc.page} page{'s' if doc.page != 1 else ''})")
        if doc.page > MAX_PAGES:
            print(f"  [WARN] {doc.page} pages - resumes should be {MAX_PAGES} pages or less (some ATS truncate after page {MAX_PAGES})")
        return output_pdf_path
    except Exception as e:
        print(f"[ERROR] Error creating PDF: {e}")
        return None

def process_experience_section(content, elements, job_style, company_style, bullet_style):
    """Process professional experience section."""
    lines = content.split('\n')
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # Look for job titles (usually contain | or dates like "2023 - Present")
        if '|' in line or re.search(r'\d{4}\s*[-–]\s*(Present|\d{4})', line):
            # This looks like a job header
            if '|' in line:
                parts = line.split('|')
                if len(parts) >= 2:
                    company = parts[0].strip().replace('&', '&amp;')
                    rest = parts[1].strip()

                    # Extract date range from the title+date string
                    date_match = _DATE_RANGE.search(rest)
                    if date_match:
                        date_str = date_match.group(1).replace('&', '&amp;')
                        title = rest[:date_match.start()].strip().replace('&', '&amp;')
                        if title:
                            job_text = f"<b>{company} | {title}</b>  <i>{date_str}</i>"
                        else:
                            job_text = f"<b>{company}</b>  <i>{date_str}</i>"
                    else:
                        rest_esc = rest.replace('&', '&amp;')
                        job_text = f"<b>{company}</b>  <i>({rest_esc})</i>"
                    elements.append(Paragraph(job_text, job_style))
            else:
                line_esc = line.replace('&', '&amp;')
                elements.append(Paragraph(f"<b>{line_esc}</b>", job_style))
            
            i += 1
            
            # Check for location line (e.g., "Remote", "Seattle, WA")
            if i < len(lines):
                next_line = lines[i].strip()
                if next_line and not next_line.startswith('•') and not ':' in next_line and len(next_line.split()) <= 3:
                    # Likely a location line
                    elements.append(Paragraph(f"<i>{next_line}</i>", company_style))
                    i += 1
            
            # Collect content for this job
            while i < len(lines):
                line = lines[i].strip()
                
                if not line:
                    i += 1
                    continue
                    
                # Check if this is the next job entry
                if '|' in line and re.search(r'\d{4}', line):
                    break
                    
                # Check for sub-sections like "Key Unity Projects:"
                if line.endswith(':') and not line.startswith('•'):
                    line_esc = line.replace('&', '&amp;')
                    elements.append(Paragraph(f"<b>{line_esc}</b>", company_style))
                    i += 1
                    
                    # Process items under this sub-section
                    while i < len(lines):
                        sub_line = lines[i].strip()
                        if not sub_line:
                            i += 1
                            break
                        if sub_line.startswith('-'):
                            # Sub-bullet with dash
                            elements.append(Paragraph('  ' + sub_line, bullet_style))
                        elif '|' in sub_line and re.search(r'\d{4}', sub_line):
                            # Next job, don't consume
                            break
                        elif sub_line.endswith(':') and not sub_line.startswith('•'):
                            # Another sub-section, don't consume
                            break
                        elif sub_line.startswith('•'):
                            # Regular bullet, end of sub-section
                            break
                        else:
                            # Part of sub-section
                            elements.append(Paragraph('  - ' + _esc(sub_line), bullet_style))
                        i += 1
                    continue
                    
                # Regular bullet points
                if line.startswith('•'):
                    elements.append(Paragraph(_esc(line), bullet_style))
                elif line and line[0].isupper() and ':' not in line:
                    # Likely a new section or standalone line
                    if re.search(r'\d{4}', line):
                        # Might be next job
                        break
                    else:
                        # Add as bullet
                        elements.append(Paragraph('• ' + line, bullet_style))
                i += 1
        else:
            i += 1

def process_skills_section(content, elements, bullet_style, normal_style):
    """Process skills section."""
    lines = content.split('\n')
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
        
        # Check if it's a categorized skill (has a colon)
        line = _esc(line)
        if ':' in line and not line.startswith('•'):
            # Bold the category
            parts = line.split(':', 1)
            formatted = f"<b>{parts[0]}:</b>{parts[1]}"
            if not line.startswith('•'):
                formatted = '• ' + formatted
            elements.append(Paragraph(formatted, bullet_style))
        elif line.startswith('•'):
            # Already a bullet point
            if ':' in line:
                parts = line.split(':', 1)
                formatted = f"{parts[0].replace('•', '• <b>')}:</b>{parts[1]}"
                elements.append(Paragraph(formatted, bullet_style))
            else:
                elements.append(Paragraph(line, bullet_style))
        else:
            # Make it a bullet point
            elements.append(Paragraph('• ' + line, bullet_style))

def process_projects_section(content, elements, project_style, bullet_style):
    """Process projects section."""
    lines = content.split('\n')
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        
        # Check if this looks like a project title (with year or not starting with bullet)
        if not line.startswith('•') and (re.search(r'\(\d{4}\)', line) or 
                                         (line[0].isupper() and not line.startswith('-'))):
            # This is a project title
            elements.append(Paragraph(f"<b>{_esc(line)}</b>", project_style))
            i += 1
            
            # Collect project description bullets
            while i < len(lines):
                desc_line = lines[i].strip()
                if not desc_line:
                    i += 1
                    break
                    
                # Check if this is a new project title
                if not desc_line.startswith('•') and not desc_line.startswith('-') and \
                   (re.search(r'\(\d{4}\)', desc_line) or 
                    (desc_line[0].isupper() and len(desc_line.split()) > 1)):
                    # Next project, don't consume
                    break
                    
                if desc_line.startswith('•'):
                    elements.append(Paragraph(_esc(desc_line), bullet_style))
                else:
                    # Add as bullet if not already
                    elements.append(Paragraph('• ' + _esc(desc_line), bullet_style))
                i += 1
        elif line.startswith('•'):
            elements.append(Paragraph(line, bullet_style))
            i += 1
        else:
            # Standalone line, treat as bullet
            elements.append(Paragraph('• ' + line, bullet_style))
            i += 1

def process_education_section(content, elements, normal_style, bullet_style):
    """Process education section."""
    lines = content.split('\n')
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
        
        # Education entries usually have degree type or school name
        if any(k in line for k in ('Master', 'Bachelor', 'University', 'School', 'Máster', 'Grado')):
            # Bold degree names
            line = re.sub(r'(Master[^,\-|]*)', r'<b>\1</b>', line)
            line = re.sub(r'(Bachelor[^,\-|]*)', r'<b>\1</b>', line)
            line = re.sub(r'((?:Máster|Grado) en [^,\-|]*)', r'<b>\1</b>', line)
            elements.append(Paragraph(line, normal_style))
        elif line.startswith('•'):
            elements.append(Paragraph(line, bullet_style))
        else:
            elements.append(Paragraph(line, normal_style))

def process_bullet_section(content, elements, bullet_style):
    """Process sections that are primarily bullet points."""
    lines = content.split('\n')
    
    for line in lines:
        line = _esc(line.strip())
        if not line:
            continue

        if not line.startswith('•'):
            line = '• ' + line

        # Bold important keywords
        line = re.sub(r'(Patent:)', r'<b>\1</b>', line)
        line = re.sub(r'(Award:)', r'<b>\1</b>', line)
        line = re.sub(r'(Published:)', r'<b>\1</b>', line)

        elements.append(Paragraph(line, bullet_style))

# Inline markup the section processors emit: <b>, <i>, <link href="...">
_MARKUP_TOKEN = re.compile(r'<(/?)(b|i|link)(?:\s+href="[^"]*")?>')


def _add_markup_runs(paragraph, markup, font_size, color, bold_default):
    """Turn ReportLab inline markup into python-docx runs."""
    from html import unescape
    from docx.shared import Pt, RGBColor

    def add_run(text, bold, italic):
        text = unescape(text)
        if not text:
            return
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Arial'
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor.from_string(color)

    bold, italic, pos = bold_default, False, 0
    for m in _MARKUP_TOKEN.finditer(markup):
        add_run(markup[pos:m.start()], bold, italic)
        closing, tag = m.group(1) == '/', m.group(2)
        if tag == 'b':
            bold = bold_default or not closing
        elif tag == 'i':
            italic = not closing
        pos = m.end()
    add_run(markup[pos:], bold, italic)


def create_resume_docx(txt_file_path, output_docx_path=None, verbose=False):
    """Create a Word .docx from a resume text file, for ATS forms that require Word.

    Reuses build_resume_elements so the .docx has the same content and structure as the PDF.
    """
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    if not os.path.exists(txt_file_path):
        print(f"Error: File {txt_file_path} not found!")
        return

    sections, elements = build_resume_elements(txt_file_path, verbose=verbose)

    if output_docx_path is None:
        output_docx_path = os.path.splitext(txt_file_path)[0] + ".docx"

    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Inches(8.5), Inches(11)
    page.left_margin = page.right_margin = Inches(0.5)
    page.top_margin = page.bottom_margin = Inches(0.5)

    name = sections['header'][0].strip().title() if sections.get('header') else 'Resume'
    document.core_properties.title = f"{name} - Resume"
    document.core_properties.author = name
    document.styles['Normal'].font.name = 'Arial'
    document.styles['Normal'].font.size = Pt(10)

    alignments = {TA_LEFT: WD_ALIGN_PARAGRAPH.LEFT, TA_CENTER: WD_ALIGN_PARAGRAPH.CENTER,
                  TA_JUSTIFY: WD_ALIGN_PARAGRAPH.JUSTIFY}
    last = None
    for el in elements:
        if isinstance(el, Spacer):
            if last is not None:
                last.paragraph_format.space_after = Pt(el.height)
            continue
        if not isinstance(el, Paragraph):
            continue
        style = el.style
        p = document.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before = Pt(style.spaceBefore)
        fmt.space_after = Pt(style.spaceAfter)
        fmt.left_indent = Pt(style.leftIndent)
        fmt.line_spacing = Pt(style.leading)
        fmt.alignment = alignments.get(style.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        color = style.textColor.hexval()[2:] if style.textColor else '000000'
        _add_markup_runs(p, el.text, style.fontSize, color, 'Bold' in style.fontName)
        last = p

    try:
        document.save(output_docx_path)
        print(f"[SUCCESS] Resume DOCX created: {output_docx_path}")
        return output_docx_path
    except Exception as e:
        print(f"[ERROR] Error creating DOCX: {e}")
        return None


def main():
    """Main function to handle command line arguments."""
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    verbose = '--verbose' in sys.argv or '-v' in sys.argv
    want_docx = '--docx' in sys.argv

    if args:
        # Process specific file
        txt_file = args[0]
        output_pdf = args[1] if len(args) > 1 else None
        create_resume_pdf(txt_file, output_pdf, verbose=verbose)
        if want_docx:
            create_resume_docx(txt_file, os.path.splitext(output_pdf or txt_file)[0] + ".docx", verbose=verbose)
    else:
        # Process all .txt resume files in the current directory
        # Look for files containing "Resume" in the name
        import glob
        resume_files = glob.glob('*Resume*.txt')

        if not resume_files:
            print("No resume .txt files found in current directory.")
            print("Usage: python create_pdf.py [resume_file.txt] [output.pdf] [--docx] [--verbose]")
            return

        for txt_file in resume_files:
            print(f"\nProcessing {txt_file}...")
            create_resume_pdf(txt_file, verbose=verbose)
            if want_docx:
                create_resume_docx(txt_file, verbose=verbose)

if __name__ == "__main__":
    main()