import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys

GREETINGS = ('Dear', 'Hi ', 'Hi,', 'Hello')


def is_body_paragraph(text):
    """A body paragraph is prose: not a greeting, and longer than a closing or signature block."""
    text = text.strip()
    return not text.startswith(GREETINGS) and len(text) > 120


def create_cover_letter_docx(txt_file_path):
    """Convert a text cover letter to DOCX format with proper formatting."""

    # Read the text content
    with open(txt_file_path, 'r', encoding='utf-8') as file:
        content = file.read()

    # Create a new Document
    doc = Document()

    # Set margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Split content into paragraphs
    paragraphs = content.split('\n\n')

    for i, para_text in enumerate(paragraphs):
        if para_text.strip():
            # Add paragraph
            paragraph = doc.add_paragraph()

            # Set font
            run = paragraph.add_run(para_text.strip())
            font = run.font
            font.name = 'Arial'
            font.size = Pt(11)

            # Justify body paragraphs; greeting, closing and signature lines stay left-aligned
            if is_body_paragraph(para_text):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            # Add spacing after paragraph (except for the last one)
            if i < len(paragraphs) - 1:
                paragraph.paragraph_format.space_after = Pt(6)

    # Save the document
    output_path = txt_file_path.replace('.txt', '.docx')
    doc.save(output_path)
    print(f"[SUCCESS] Cover letter DOCX created: {output_path}")
    return output_path

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python create_cover_letter_docx.py <cover_letter.txt>")
        sys.exit(1)
    txt_path = sys.argv[1]

    if os.path.exists(txt_path):
        create_cover_letter_docx(txt_path)
    else:
        print(f"[ERROR] File not found: {txt_path}")